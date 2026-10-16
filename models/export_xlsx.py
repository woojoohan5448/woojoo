"""Excel export of the full intelligence dataset. python3 -m models.export_xlsx"""
from __future__ import annotations

import json

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from models.common import EXPORTS, PROCESSED

HEAD = PatternFill("solid", fgColor="1F1F24")


def _sheet(wb, title, header, rows, widths=None):
    ws = wb.create_sheet(title)
    ws.append(header)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = HEAD
        c.alignment = Alignment(vertical="center", wrap_text=True)
    for r in rows:
        ws.append(r)
    ws.freeze_panes = "B2"
    for i, w in enumerate(widths or [], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    return ws


def v(r):
    return None if r is None or r.get("value") is None else r["value"]


def export(path=None):
    d = json.loads((PROCESSED / "dashboard.json").read_text(encoding="utf-8"))
    wb = Workbook()
    wb.remove(wb.active)
    readme = wb.create_sheet("README")
    for line in [
        "LG U+ Enterprise — 53 Key Account IT Spending & Full Potential Revenue (2026–2031)",
        f"data_version {d['meta']['data_version']} · 단위 억원(KRW) · 생성 {d['meta']['generated_at']}",
        "상태: VERIFIED(원천 직접 확인) / DERIVED(확정 산식) / ESTIMATED(추정) / SCENARIO(시나리오) / UNKNOWN(미확인 — 0 아님)",
        "CR(Current Revenue)·Target은 사내 입력 전까지 NULL. 합계는 근거 있는 고객 부분합(Coverage 열 참조).",
        "신뢰도 = 분석 품질 평가 점수(0–100), 확률 아님. 상세는 docs/METHODOLOGY.md.",
    ]:
        readme.append([line])
    readme.column_dimensions["A"].width = 120

    rows = []
    for a in d["accounts"]:
        m, b, t = a["m"]["2026"], a["m"]["2031"], a["m"]["target"]
        rows.append([a["account_id"], a["original_name"], a["entity"].get("legal_name_ko"), a["entity"].get("corp_code"), a["segment"],
                     a["research_status"], a["aggregate_class"], a["entity"].get("entity_status"), a["anchor"]["method"], a["anchor"].get("fiscal_year"),
                     v(m["TS"]), m["TS"]["low"], m["TS"]["high"], m["TS"]["status"], m["TS"]["confidence"], v(m["RAW"]), v(m["FPR"]), v(m["CR"]),
                     v(m["CSOW"]), v(m["PCR"]), v(m["UPSIDE"]),
                     v(b["low"]["TS"]), v(b["base"]["TS"]), v(b["high"]["TS"]), v(b["base"]["RAW_base"]), v(b["base"]["RAW_expansion"]),
                     v(b["low"]["FPR_base"]), v(b["base"]["FPR_base"]), v(b["high"]["FPR_base"]), v(b["base"]["FPR_expansion"]),
                     v(t["TARGET"]), t["TARGET"]["status"], v(t["TARGET_SOW"]), v(b["base"]["TS_CAGR"]), v(b["base"]["RAW_CAGR"]), v(b["base"]["FPR_CAGR"])])
    _sheet(wb, "Accounts", ["ID", "원본 고객명", "법인명", "DART corp_code", "산업", "조사상태", "합산분류", "법인매핑", "TS 앵커", "앵커 FY",
                            "TS 2026", "TS 2026 Low", "TS 2026 High", "TS 상태", "TS 신뢰도", "RAW 2026", "FPR 2026", "CR 2026", "C-SOW", "PCR", "Upside 2026",
                            "TS 2031 Low", "TS 2031 Base", "TS 2031 High", "RAW 2031 Base", "RAW 2031 Expansion", "FPR 2031 Low", "FPR 2031 Base", "FPR 2031 High",
                            "FPR 2031 Expansion", "Target 2031", "Target 상태", "Target SOW", "TS CAGR", "RAW CAGR", "FPR CAGR"], rows,
           [6, 22, 26, 12, 14, 16, 18, 11, 20, 8] + [12] * 26)

    lineage = []
    for a in d["accounts"]:
        for k, r in a["m"]["2026"].items():
            lineage.append([a["account_id"], a["original_name"], k, 2026, "base", v(r), r["status"], r["confidence"], r["formula"],
                            ", ".join(r["evidence_ids"]), ", ".join(r["assumption_ids"])])
    _sheet(wb, "Lineage_2026", ["ID", "고객", "지표", "연도", "시나리오", "값", "상태", "신뢰도", "산식", "출처 ID", "가정 ID"], lineage, [6, 20, 8, 6, 8, 12, 11, 8, 70, 30, 20])

    ev = [[e.get("evidence_id"), e.get("account_id"), e.get("tier"), e.get("source_type"), e.get("source_title"), e.get("publisher"), e.get("source_url"),
           e.get("document_id"), e.get("publication_date"), e.get("accessed_at"), e.get("fiscal_period"), e.get("page_or_section"),
           e.get("quoted_statement"), e.get("reported_value"), e.get("reported_unit"), e.get("verification_status"), e.get("source_hash")]
          for a in d["accounts"] for e in a["evidence"]]
    _sheet(wb, "Evidence", ["evidence_id", "account_id", "tier", "source_type", "제목", "발행처", "URL", "문서ID", "발행일", "조회일", "회계기간", "위치",
                            "원문 인용", "원천값", "단위", "검증상태", "hash"], ev, [12, 8, 5, 18, 40, 16, 40, 18, 11, 11, 10, 24, 60, 12, 8, 20, 18])
    iss = [[i["account_id"], i["severity"], i["code"], i.get("metric"), i.get("year"), i.get("scenario"), i["message"]] for a in d["accounts"] for i in a["issues"]]
    _sheet(wb, "Data_Quality", ["account_id", "심각도", "코드", "지표", "연도", "시나리오", "내용"], iss, [8, 8, 22, 10, 6, 8, 100])
    st = [[s["area"], s["priority_score"], s["FPR_2026"], s["FPR_2031"], s["FPR_2031_expansion"], s["FPR_CAGR"], s["n_accounts"], s["demand_confidence"],
           s.get("entry_feasibility"), s.get("competition"), s.get("capability_gap"), s.get("investment_need"), s.get("difficulty"), s.get("recommendation_type"),
           s.get("recommendation"), s.get("caveats")] for s in d["strategy"]]
    _sheet(wb, "Strategy", ["사업영역", "우선순위 점수", "FPR 2026", "FPR 2031", "FPR 2031 Expansion", "FPR CAGR", "수요 고객 수", "수요 신뢰도", "진입 가능성",
                            "경쟁", "역량 격차", "투자 필요", "난이도", "권고 유형", "권고", "주의"], st, [22, 10, 12, 12, 14, 9, 9, 9, 10, 40, 14, 12, 10, 12, 70, 50])
    gaps = [[a["account_id"], a["original_name"], a["research_status"], a["anchor"]["method"], "; ".join(a.get("open_questions", [])[:6]),
             "; ".join(a.get("override_notes", []))] for a in d["accounts"] if a["m"]["2026"]["TS"]["value"] is None or a["m"]["2026"]["TS"]["confidence"] < 35]
    _sheet(wb, "Research_Gaps", ["ID", "고객", "조사상태", "TS 방법", "미확인 항목", "비고"], gaps, [6, 22, 16, 18, 100, 80])
    for ws in wb.worksheets[1:]:
        for row in ws.iter_rows(min_row=2):
            for c in row:
                if isinstance(c.value, float):
                    c.number_format = "#,##0.00" if abs(c.value) < 1 else "#,##0"
    EXPORTS.mkdir(parents=True, exist_ok=True)
    out = path or (EXPORTS / "key_account_intelligence.xlsx")
    wb.save(out)
    return out


if __name__ == "__main__":
    print(export())
