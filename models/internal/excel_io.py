"""Phase 3 — Internal Excel Import Engine (사내 데이터 전용).

Workflow (docs/INTERNAL_EXCEL_SPEC.md):
 1 파일 선택 → 2 형식·컬럼 검사 → 3 고객 매핑 → 4 연도·단위·통화 검증 → 5 기존값 비교 → 6 차이·충돌 보고
 → 7 사용자 검토·승인 → 8 승인 데이터 저장(버전) → 9 재무모델 재계산 → 10 대시보드 갱신

CLI:
  python3 -m models.internal.excel_io template                    # 빈 입력 템플릿 생성
  python3 -m models.internal.excel_io preview <file.xlsx>          # 1~6 단계, 배치 생성(PENDING_REVIEW)
  python3 -m models.internal.excel_io approve <batch_id> --approver 홍길동 [--rows 1,2,5]   # 7~8
  python3 -m models.internal.excel_io reject <batch_id> --approver 홍길동
  python3 -m models.internal.excel_io status

보안: 이 모듈은 네트워크 호출을 하지 않는다. 저장 위치(INTERNAL_STORE_DIR, 기본 data/internal_store)는 .gitignore 대상.
외부 LLM/MCP로 사내 데이터를 전송하는 코드 경로가 없다.
"""
from __future__ import annotations

import csv
import difflib
import hashlib
import json
import os
import re
import sys
import uuid
from datetime import datetime, timezone
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parents[2]
STORE = Path(os.environ.get("INTERNAL_STORE_DIR", ROOT / "data" / "internal_store"))
TEMPLATE_DIR = ROOT / "data" / "internal_inputs" / "templates"

LAYERS = ("INTERNAL_ACTUAL", "SALES_ESTIMATE", "MANAGEMENT_TARGET", "APPROVED_ACCOUNT_PLAN")
UNIT_TO_EOK = {"원": Decimal("1e-8"), "천원": Decimal("1e-5"), "백만원": Decimal("0.01"), "억원": Decimal(1), "조원": Decimal(10000)}

SHEETS = {
    "INTERNAL_ACTUAL": {"layer": "INTERNAL_ACTUAL", "required": ["account", "fiscal_year", "period_type", "revenue", "unit"],
                        "optional": ["as_of", "product_line", "contract_id", "source_system", "note"],
                        "enums": {"period_type": ["YTD", "FORECAST", "ACTUAL_FULL_YEAR"]},
                        "desc": "LG U+ 고객별 실제 매출. 2026 YTD(기준일)와 2026 연말 예상(FORECAST) 구분. 상품/계약 단위 행은 원래 정의 그대로 유지(합산은 엔진이 수행)."},
    "SALES_ESTIMATE": {"layer": "SALES_ESTIMATE", "required": ["account", "fiscal_year", "metric", "value", "unit", "sales_rep"],
                       "optional": ["as_of", "rationale", "confidence"],
                       "enums": {"metric": ["TS", "RAW", "FPR"]},
                       "desc": "영업사원 추정 TS/RAW/FPR (2026, 2031). 외부 추정과 비교(Reconciliation)용 — 승인 없이 최종값을 덮어쓰지 않음."},
    "MANAGEMENT_TARGET": {"layer": "MANAGEMENT_TARGET", "required": ["account", "fiscal_year", "target_revenue", "unit", "approver"],
                          "optional": ["as_of", "basis"], "enums": {},
                          "desc": "경영진 승인 2031 목표 매출 (Target ≤ FPR_2031 검증)."},
    "ACCOUNT_PLAN": {"layer": "APPROVED_ACCOUNT_PLAN", "required": ["account", "assumption_key", "value"],
                     "optional": ["unit", "fiscal_year", "owner", "note"], "enums": {},
                     "desc": "Account Plan 내부 가정(예: 재계약 시점, 회선 점유율, 상면 kW). 키-값 형태."},
}


# ---------------------------------------------------------------- account mapping
def _norm(s: str) -> str:
    s = (s or "").lower()
    s = re.sub(r"\(주\)|㈜|주식회사|유한회사|유한책임회사|co\.?,?\s*ltd\.?|inc\.?|corporation|corp\.?|korea|코리아|\s|[.,()\-·]", "", s)
    return s


def build_alias_index() -> tuple[dict, list[dict]]:
    master = list(csv.DictReader(open(ROOT / "docs" / "ACCOUNT_MASTER.csv", encoding="utf-8-sig")))
    idx: dict[str, set] = {}

    def put(alias, aid):
        if alias:
            idx.setdefault(_norm(alias), set()).add(aid)
    for r in master:
        aid = r["account_id"]
        for k in ("account_id", "original_name", "legal_name_ko", "legal_name_en", "corp_code", "stock_code"):
            put(r.get(k), aid)
    extra = ROOT / "config" / "account_aliases.json"
    if extra.exists():
        for aid, aliases in json.loads(extra.read_text(encoding="utf-8")).items():
            if aid.startswith("_"):
                continue
            for al in aliases:
                put(al, aid)
    return idx, master


def map_account(name: str, idx: dict) -> dict:
    n = _norm(str(name))
    if not n:
        return {"status": "UNMAPPED", "account_id": None, "method": "empty"}
    if n in idx:
        ids = idx[n]
        return {"status": "MAPPED", "account_id": next(iter(ids)), "method": "exact_alias"} if len(ids) == 1 else \
            {"status": "AMBIGUOUS", "account_id": None, "candidates": sorted(ids), "method": "exact_alias"}
    scored = sorted(((difflib.SequenceMatcher(None, n, k).ratio(), k) for k in idx), reverse=True)[:3]
    if scored and scored[0][0] >= 0.88 and (len(scored) == 1 or scored[0][0] - scored[1][0] >= 0.05) and len(idx[scored[0][1]]) == 1:
        return {"status": "MAPPED", "account_id": next(iter(idx[scored[0][1]])), "method": f"fuzzy({scored[0][0]:.2f})", "needs_review": True}
    return {"status": "UNMAPPED" if not scored or scored[0][0] < 0.6 else "AMBIGUOUS", "account_id": None,
            "candidates": [sorted(idx[k])[0] for _, k in scored], "method": "fuzzy_low"}


# ---------------------------------------------------------------- template
def make_template(path: Path | None = None) -> Path:
    _, master = build_alias_index()
    wb = Workbook()
    ws = wb.active; ws.title = "README"
    for line in ["LG U+ Enterprise Key Account — 사내 데이터 입력 템플릿 (v1)",
                 "⚠ 사내 기밀. 외부 LLM·외부 MCP·외부 저장소로 전송 금지. 승인된 사내 환경에서만 처리.",
                 "account 열: Account ID(KA01…) 또는 고객명. 고객명은 별칭 규칙으로 자동 매핑되며 모호하면 검토 대상이 됩니다.",
                 "unit 열: 원 / 천원 / 백만원 / 억원 / 조원 (엔진이 억원으로 환산). 통화는 KRW만 허용.",
                 "fiscal_year: 2025~2031. INTERNAL_ACTUAL.period_type: YTD(기준일 as_of 필수) / FORECAST(연말 예상) / ACTUAL_FULL_YEAR.",
                 "업로드 후 preview → 차이 보고 검토 → approve 시에만 저장·재계산됩니다. 기존 값은 버전 이력으로 보존됩니다."]:
        ws.append([line])
    ws.column_dimensions["A"].width = 120
    head = PatternFill("solid", fgColor="1F1F24")
    for sheet, spec in SHEETS.items():
        s = wb.create_sheet(sheet)
        cols = spec["required"] + spec["optional"]
        s.append(cols)
        for c in s[1]:
            c.font = Font(bold=True, color="FFFFFF"); c.fill = head
        s.append([f"설명: {spec['desc']}"] + [""] * (len(cols) - 1))
        s["A2"].font = Font(italic=True, color="888888")
        for col, vals in spec["enums"].items():
            letter = chr(ord("A") + cols.index(col))
            dv = DataValidation(type="list", formula1='"' + ",".join(vals) + '"', allow_blank=True)
            s.add_data_validation(dv); dv.add(f"{letter}3:{letter}2000")
        unit_letter = chr(ord("A") + cols.index("unit")) if "unit" in cols else None
        if unit_letter:
            dv = DataValidation(type="list", formula1='"원,천원,백만원,억원,조원"', allow_blank=True)
            s.add_data_validation(dv); dv.add(f"{unit_letter}3:{unit_letter}2000")
        for i, _ in enumerate(cols):
            s.column_dimensions[chr(ord("A") + i)].width = 18
    ref = wb.create_sheet("ACCOUNTS_REF")
    ref.append(["account_id", "original_name", "legal_name_ko", "legal_name_en", "corp_code"])
    for r in master:
        ref.append([r["account_id"], r["original_name"], r["legal_name_ko"], r["legal_name_en"], r["corp_code"]])
    TEMPLATE_DIR.mkdir(parents=True, exist_ok=True)
    path = path or TEMPLATE_DIR / "KeyAccount_Internal_Input_Template.xlsx"
    wb.save(path)
    return path


# ---------------------------------------------------------------- store (versioned)
def _store_path(layer: str) -> Path:
    return STORE / f"{layer}.json"


def load_layer(layer: str) -> dict:
    p = _store_path(layer)
    return json.loads(p.read_text(encoding="utf-8")) if p.exists() else {"layer": layer, "current": {}, "history": []}


def current_values() -> dict:
    return {l: load_layer(l)["current"] for l in LAYERS}


def _key(layer: str, r: dict) -> str:
    if layer == "INTERNAL_ACTUAL":
        return f"{r['account_id']}|{r['fiscal_year']}|{r['period_type']}|{r.get('product_line') or 'ALL'}|{r.get('contract_id') or 'ALL'}"
    if layer == "SALES_ESTIMATE":
        return f"{r['account_id']}|{r['fiscal_year']}|{r['metric']}"
    if layer == "MANAGEMENT_TARGET":
        return f"{r['account_id']}|{r['fiscal_year']}"
    return f"{r['account_id']}|{r['assumption_key']}"


# ---------------------------------------------------------------- preview (validate + diff)
def _num(v):
    if v is None or (isinstance(v, str) and not v.strip()):
        return None
    try:
        return Decimal(str(v).replace(",", "").strip())
    except InvalidOperation:
        raise ValueError(f"숫자 아님: {v}")


def preview(path: str | Path) -> dict:
    path = Path(path)
    idx, master = build_alias_index()
    names = {r["account_id"]: r["original_name"] for r in master}
    wb = load_workbook(path, data_only=True)
    batch = {"batch_id": datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%SZ") + "-" + uuid.uuid4().hex[:6],
             "file": path.name, "file_sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
             "created_at": datetime.now(timezone.utc).isoformat(), "status": "PENDING_REVIEW", "rows": [], "issues": [], "summary": {}}
    cur = current_values()
    rid = 0
    for sheet, spec in SHEETS.items():
        if sheet not in wb.sheetnames:
            continue
        ws = wb[sheet]
        header = [str(c.value).strip() if c.value is not None else "" for c in ws[1]]
        missing = [c for c in spec["required"] if c not in header]
        if missing:
            batch["issues"].append({"sheet": sheet, "row": 1, "severity": "ERROR", "code": "MISSING_COLUMNS", "message": f"필수 컬럼 누락: {missing}"})
            continue
        for ri, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            rec = {h: row[i] if i < len(row) else None for i, h in enumerate(header) if h}
            if not any(v not in (None, "") for v in rec.values()) or str(rec.get(header[0]) or "").startswith("설명:"):
                continue
            rid += 1
            item = {"row_id": rid, "sheet": sheet, "excel_row": ri, "layer": spec["layer"], "raw": {k: (str(v) if v is not None else None) for k, v in rec.items()},
                    "errors": [], "warnings": []}
            m = map_account(rec.get("account"), idx)
            item["mapping"] = m
            if m["status"] != "MAPPED":
                item["errors"].append(f"고객 매핑 {m['status']}: '{rec.get('account')}' 후보 {m.get('candidates')}")
            elif m.get("needs_review"):
                item["warnings"].append(f"유사도 매핑 {m['method']} → {m['account_id']} {names.get(m['account_id'])} (확인 필요)")
            for col, vals in spec["enums"].items():
                if rec.get(col) not in vals:
                    item["errors"].append(f"{col} 값 '{rec.get(col)}' 허용값 아님 {vals}")
            fy = rec.get("fiscal_year")
            if fy is not None:
                try:
                    fy = int(fy)
                    if not 2025 <= fy <= 2031:
                        item["errors"].append(f"fiscal_year {fy} 범위(2025~2031) 밖")
                except ValueError:
                    item["errors"].append(f"fiscal_year 형식 오류: {fy}")
            elif "fiscal_year" in spec["required"]:
                item["errors"].append("fiscal_year 누락")
            val_col = {"INTERNAL_ACTUAL": "revenue", "SALES_ESTIMATE": "value", "MANAGEMENT_TARGET": "target_revenue"}.get(spec["layer"])
            value_eok = None
            if val_col:
                try:
                    v = _num(rec.get(val_col))
                    unit = (rec.get("unit") or "").strip()
                    if unit not in UNIT_TO_EOK:
                        item["errors"].append(f"단위 '{unit}' 미지원(원/천원/백만원/억원/조원)")
                    elif v is None:
                        item["errors"].append(f"{val_col} 값 없음")
                    elif v < 0:
                        item["errors"].append(f"{val_col} 음수")
                    else:
                        value_eok = v * UNIT_TO_EOK[unit]
                        if value_eok > 100000:
                            item["warnings"].append(f"값 {value_eok:,.0f}억 — 단위 오기 의심")
                except ValueError as e:
                    item["errors"].append(str(e))
            if spec["layer"] == "INTERNAL_ACTUAL" and rec.get("period_type") == "YTD" and not rec.get("as_of"):
                item["errors"].append("YTD는 as_of(기준일) 필수")
            if spec["layer"] == "MANAGEMENT_TARGET" and fy not in (None, 2031):
                item["warnings"].append("목표는 2031 기준 — 다른 연도는 참고로만 저장")
            norm = {"account_id": m.get("account_id"), "fiscal_year": fy if isinstance(fy, int) else None,
                    **{k: rec.get(k) for k in spec["required"] + spec["optional"] if k not in ("account", "fiscal_year")}}
            if value_eok is not None:
                norm["value_eok"] = str(value_eok)
            item["normalized"] = norm
            if not item["errors"]:
                k = _key(spec["layer"], norm)
                old = cur[spec["layer"]].get(k)
                item["key"] = k
                item["diff"] = {"type": "NEW" if old is None else ("UNCHANGED" if old.get("value_eok") == norm.get("value_eok") and old.get("value") == norm.get("value") else "CHANGED"),
                                "old_value_eok": (old or {}).get("value_eok"), "new_value_eok": norm.get("value_eok")}
            batch["rows"].append(item)
    # duplicate keys inside file
    seen = {}
    for it in batch["rows"]:
        if it.get("key"):
            if it["key"] in seen:
                it["errors"].append(f"파일 내 중복 키 (행 {seen[it['key']]})")
            seen[it["key"]] = it["excel_row"]
    from collections import Counter
    batch["summary"] = {"rows": len(batch["rows"]), "valid": sum(not r["errors"] for r in batch["rows"]),
                        "invalid": sum(bool(r["errors"]) for r in batch["rows"]),
                        "diff": dict(Counter(r.get("diff", {}).get("type", "INVALID") for r in batch["rows"])),
                        "by_layer": dict(Counter(r["layer"] for r in batch["rows"]))}
    _save_batch(batch)
    return batch


def _batch_dir() -> Path:
    d = STORE / "batches"
    d.mkdir(parents=True, exist_ok=True)
    return d


def _save_batch(b):
    (_batch_dir() / f"{b['batch_id']}.json").write_text(json.dumps(b, ensure_ascii=False, indent=1), encoding="utf-8")


def load_batch(batch_id: str) -> dict:
    return json.loads((_batch_dir() / f"{batch_id}.json").read_text(encoding="utf-8"))


def approve(batch_id: str, approver: str, rows: list[int] | None = None) -> dict:
    if not approver:
        raise ValueError("approver 필수")
    b = load_batch(batch_id)
    if b["status"] != "PENDING_REVIEW":
        raise ValueError(f"배치 상태 {b['status']} — 승인 불가")
    now = datetime.now(timezone.utc).isoformat()
    applied = 0
    layers = {l: load_layer(l) for l in LAYERS}
    for it in b["rows"]:
        if it["errors"] or (rows and it["row_id"] not in rows) or it["diff"]["type"] == "UNCHANGED":
            continue
        L = layers[it["layer"]]
        k = it["key"]
        if k in L["current"]:
            L["history"].append({**L["current"][k], "superseded_at": now, "superseded_by_batch": batch_id})
        L["current"][k] = {**it["normalized"], "batch_id": batch_id, "approved_by": approver, "approved_at": now,
                           "source_file": b["file"], "source_sha256": b["file_sha256"], "excel_row": it["excel_row"]}
        applied += 1
    STORE.mkdir(parents=True, exist_ok=True)
    for l, L in layers.items():
        _store_path(l).write_text(json.dumps(L, ensure_ascii=False, indent=1), encoding="utf-8")
    b.update({"status": "APPROVED", "approved_by": approver, "approved_at": now, "applied_rows": applied, "approved_row_ids": rows})
    _save_batch(b)
    return b


def reject(batch_id: str, approver: str) -> dict:
    b = load_batch(batch_id)
    b.update({"status": "REJECTED", "rejected_by": approver, "rejected_at": datetime.now(timezone.utc).isoformat()})
    _save_batch(b)
    return b


def management_rows_for_build() -> dict:
    """Approved internal layers → per-account inputs consumed by models.build (CR, Target, sales estimates)."""
    cur = current_values()
    out: dict[str, dict] = {}
    for k, r in cur["INTERNAL_ACTUAL"].items():
        aid, fy, pt = r["account_id"], r["fiscal_year"], r["period_type"]
        o = out.setdefault(aid, {"account_id": aid})
        if fy == 2026 and pt in ("FORECAST", "ACTUAL_FULL_YEAR"):
            o["cr_2026_forecast_eok"] = str(Decimal(o.get("cr_2026_forecast_eok") or 0) + Decimal(r["value_eok"]))
            o["cr_basis"] = f"INTERNAL_ACTUAL 승인 배치 {r['batch_id']} ({r['approved_by']})"
        if fy == 2026 and pt == "YTD":
            o["cr_2026_ytd_eok"] = str(Decimal(o.get("cr_2026_ytd_eok") or 0) + Decimal(r["value_eok"]))
            o["cr_2026_ytd_asof"] = r.get("as_of")
    for k, r in cur["MANAGEMENT_TARGET"].items():
        if r["fiscal_year"] == 2031:
            o = out.setdefault(r["account_id"], {"account_id": r["account_id"]})
            o["target_revenue_2031_eok"] = r["value_eok"]
            o["target_basis"] = f"MANAGEMENT_TARGET 승인 ({r['approved_by']}, {r.get('basis') or ''})"
    for k, r in cur["SALES_ESTIMATE"].items():
        o = out.setdefault(r["account_id"], {"account_id": r["account_id"]})
        o.setdefault("sales_estimates", {})[f"{r['fiscal_year']}|{r['metric']}"] = {"value_eok": r["value_eok"], "sales_rep": r.get("sales_rep"), "as_of": r.get("as_of"), "rationale": r.get("rationale")}
    return out


def status():
    rows = []
    for p in sorted(_batch_dir().glob("*.json")):
        b = json.loads(p.read_text(encoding="utf-8"))
        rows.append({k: b.get(k) for k in ("batch_id", "file", "status", "created_at", "approved_by", "applied_rows")} | {"summary": b.get("summary")})
    return {"store": str(STORE), "layers": {l: len(load_layer(l)["current"]) for l in LAYERS}, "batches": rows}


if __name__ == "__main__":
    cmd = sys.argv[1] if len(sys.argv) > 1 else "status"
    arg = lambda k: sys.argv[sys.argv.index(k) + 1] if k in sys.argv else None
    if cmd == "template":
        print(make_template())
    elif cmd == "preview":
        b = preview(sys.argv[2])
        print(json.dumps({"batch_id": b["batch_id"], "summary": b["summary"], "issues": b["issues"],
                          "row_errors": [(r["excel_row"], r["sheet"], r["errors"]) for r in b["rows"] if r["errors"]][:30]}, ensure_ascii=False, indent=1))
    elif cmd == "approve":
        rws = [int(x) for x in arg("--rows").split(",")] if arg("--rows") else None
        b = approve(sys.argv[2], arg("--approver"), rws)
        print(f"APPROVED {b['batch_id']} applied={b['applied_rows']} → 다음: python3 -m models.build")
    elif cmd == "reject":
        print(reject(sys.argv[2], arg("--approver"))["status"])
    else:
        print(json.dumps(status(), ensure_ascii=False, indent=1))
