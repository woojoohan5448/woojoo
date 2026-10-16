"""TEST-ONLY QA re-verification of the internal Excel integration (synthetic values, sandbox only).
Run via tests/test_qa_integration.py which sets INTERNAL_STORE_DIR / KAI_DATA_OUT to a temp dir.

Checks: layer separation, external estimates never overwritten, conflict → review flags,
pending/rejected batches do not touch the store, FORECAST vs ACTUAL_FULL_YEAR not double-counted,
explicit empty row selection approves nothing, approved batch cannot be rejected.
"""
import json
import os
import sys
from decimal import Decimal
from pathlib import Path

assert os.environ.get("INTERNAL_STORE_DIR") and os.environ.get("KAI_DATA_OUT"), "sandbox env required"
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from openpyxl import load_workbook  # noqa: E402

from models.build import main as build  # noqa: E402
from models.internal import excel_io as X  # noqa: E402

tmp = Path(os.environ["KAI_DATA_OUT"])
res = {}
EXT = [("2026", "TS"), ("2026", "RAW"), ("2026", "FPR")]


def ext_values(out, aids):
    acc = {a["account_id"]: a for a in out["accounts"]}
    v = {}
    for aid in aids:
        a = acc[aid]
        for y, m in EXT:
            v[f"{aid}|{y}|{m}"] = str(a["m"][y][m]["value"])
        b = a["m"]["2031"]["base"]
        for k in ("TS", "RAW_base", "FPR_base"):
            v[f"{aid}|2031|{k}"] = str(b[k]["value"])
    return v


def workbook(path, actual=(), sales=(), target=(), plan=()):
    wb = load_workbook(X.make_template(tmp / "tpl.xlsx"))
    for sheet, rows in (("INTERNAL_ACTUAL", actual), ("SALES_ESTIMATE", sales), ("MANAGEMENT_TARGET", target), ("ACCOUNT_PLAN", plan)):
        ws = wb[sheet]
        hdr = [c.value for c in ws[1]]
        for r in rows:
            ws.append([r.get(h) for h in hdr])
    wb.save(path)
    return path


# 0) baseline without internal data
base = build()
AIDS = ["KA03", "KA18", "KA09"]
ext0 = ext_values(base, AIDS)
fpr = {a["account_id"]: a["m"]["2026"]["FPR"]["value"] for a in base["accounts"]}
fpr31 = {a["account_id"]: a["m"]["2031"]["base"]["FPR_base"]["value"] for a in base["accounts"]}
raw26 = {a["account_id"]: a["m"]["2026"]["RAW"]["value"] for a in base["accounts"]}
res["baseline_internal"] = base["meta"]["contains_internal_data"]

f1 = workbook(tmp / "QA_upload1.xlsx",
              actual=[{"account": "KA03", "fiscal_year": 2026, "period_type": "FORECAST", "revenue": 100, "unit": "억원", "note": "TEST-ONLY"},
                      {"account": "KA18", "fiscal_year": 2026, "period_type": "FORECAST", "revenue": float(fpr["KA18"]) * 3, "unit": "억원", "note": "TEST-ONLY CR>FPR"}],
              sales=[{"account": "KA03", "fiscal_year": 2026, "metric": "TS", "value": 3000, "unit": "억원", "sales_rep": "TEST-REP"},
                     {"account": "KA03", "fiscal_year": 2026, "metric": "RAW", "value": round(float(raw26["KA03"]) * 1.05, 2), "unit": "억원", "sales_rep": "TEST-REP"},
                     {"account": "KA03", "fiscal_year": 2026, "metric": "FPR", "value": round(float(fpr["KA03"]) * 1.6, 2), "unit": "억원", "sales_rep": "TEST-REP"}],
              target=[{"account": "KA03", "fiscal_year": 2031, "target_revenue": round(float(fpr31["KA03"]) * 2, 2), "unit": "억원", "approver": "TEST-MGMT"},
                      {"account": "KA18", "fiscal_year": 2031, "target_revenue": 100, "unit": "억원", "approver": "TEST-MGMT"}],
              plan=[{"account": "KA03", "assumption_key": "connectivity_share", "value": 0.5, "owner": "TEST"}])

# 1) preview only → store untouched
b1 = X.preview(f1)
res["pending_status"] = b1["status"]
res["layers_after_preview"] = X.status()["layers"]

# 2) explicit empty selection approves nothing
e = X.approve(b1["batch_id"], "TEST-APPROVER", rows=[])
res["empty_selection_applied"] = e["applied_rows"]
res["layers_after_empty"] = X.status()["layers"]

# 3) rejected batch leaves store untouched; approved batch cannot be rejected
b_rej = X.preview(f1)
X.reject(b_rej["batch_id"], "TEST-APPROVER")
res["layers_after_reject"] = X.status()["layers"]
b2 = X.preview(f1)
X.approve(b2["batch_id"], "TEST-APPROVER")
try:
    X.reject(b2["batch_id"], "TEST-APPROVER")
    res["reject_approved_blocked"] = False
except ValueError:
    res["reject_approved_blocked"] = True
res["layers_after_approve"] = X.status()["layers"]

# 4) layer separation: each layer holds only its own record kinds
cur = X.current_values()
res["layer_keys"] = {l: sorted(cur[l].keys()) for l in X.LAYERS}
res["actual_has_metric_field"] = any("metric" in r for r in cur["INTERNAL_ACTUAL"].values())
res["sales_has_period_type"] = any("period_type" in r for r in cur["SALES_ESTIMATE"].values())
res["target_layer_accounts"] = sorted(r["account_id"] for r in cur["MANAGEMENT_TARGET"].values())

out = build()
acc = {a["account_id"]: a for a in out["accounts"]}
ext1 = ext_values(out, AIDS)
res["external_unchanged"] = ext1 == ext0
res["external_diff"] = {k: (ext0[k], ext1[k]) for k in ext0 if ext0[k] != ext1[k]}
kb, sh = acc["KA03"], acc["KA18"]
res["kb_cr"] = str(kb["m"]["2026"]["CR"]["value"])
res["kb_flags"] = {f"{r['year']}|{r['metric']}": r["flag"] for r in kb["reconciliation"]["rows"]}
res["kb_target_kept"] = str(kb["m"]["target"]["TARGET"]["value"])
res["kb_target_error"] = any(i["code"] == "TARGET_GT_FPR" and i["severity"] == "ERROR" for i in kb["issues"])
res["sh_cr_kept"] = str(sh["m"]["2026"]["CR"]["value"])
res["sh_cr_gt_fpr_warn"] = any(i["code"] == "CR_GT_FPR" for i in sh["issues"])
res["sh_recon_note"] = sh["reconciliation"]["notes"]
res["ka09_no_recon"] = acc["KA09"]["reconciliation"] is None
res["ka09_cr_null"] = acc["KA09"]["m"]["2026"]["CR"]["value"] is None
res["plan_layer_stored"] = len(cur["APPROVED_ACCOUNT_PLAN"])
res["contains_internal"] = out["meta"]["contains_internal_data"]

# 5) FORECAST + ACTUAL_FULL_YEAR for the same account/year must not be summed
f3 = workbook(tmp / "QA_upload3.xlsx",
              actual=[{"account": "KA03", "fiscal_year": 2026, "period_type": "ACTUAL_FULL_YEAR", "revenue": 120, "unit": "억원"}])
b3 = X.preview(f3)
X.approve(b3["batch_id"], "TEST-APPROVER")
out3 = build()
res["kb_cr_after_actual"] = str(next(a for a in out3["accounts"] if a["account_id"] == "KA03")["m"]["2026"]["CR"]["value"])
res["history_len_actual"] = len(X.load_layer("INTERNAL_ACTUAL")["history"])
print("RESULT_JSON=" + json.dumps(res, ensure_ascii=False, default=str))
