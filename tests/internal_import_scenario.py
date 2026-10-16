"""TEST-ONLY end-to-end scenario for Phase 3/4 (synthetic, clearly fake internal values).
Run only via tests/test_internal_import.py, which sets INTERNAL_STORE_DIR and KAI_DATA_OUT to a temp dir
so nothing touches production data/internal_store or data/processed.
"""
import json
import os
import sys
from decimal import Decimal
from pathlib import Path

assert os.environ.get("INTERNAL_STORE_DIR") and os.environ.get("KAI_DATA_OUT"), "sandbox env required"
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from openpyxl import load_workbook  # noqa: E402

from models.internal import excel_io as X  # noqa: E402

tmp = Path(os.environ["KAI_DATA_OUT"])
res = {}


def workbook(path, actual, sales, target):
    t = X.make_template(tmp / "tpl.xlsx")
    wb = load_workbook(t)
    for sheet, rows in (("INTERNAL_ACTUAL", actual), ("SALES_ESTIMATE", sales), ("MANAGEMENT_TARGET", target)):
        ws = wb[sheet]
        hdr = [c.value for c in ws[1]]
        for r in rows:
            ws.append([r.get(h) for h in hdr])
    wb.save(path)
    return path


# 1) first upload — name variants, unit conversion, a bad unit, an ambiguous name
f1 = workbook(tmp / "TEST_upload1.xlsx",
              actual=[{"account": "국민 은행", "fiscal_year": 2026, "period_type": "FORECAST", "revenue": 20000, "unit": "백만원", "note": "TEST-ONLY"},
                      {"account": "KA03", "fiscal_year": 2026, "period_type": "YTD", "revenue": 150, "unit": "억원", "as_of": "2026-09-30"},
                      {"account": "카카오", "fiscal_year": 2026, "period_type": "FORECAST", "revenue": 50, "unit": "억원"},
                      {"account": "KA09", "fiscal_year": 2026, "period_type": "FORECAST", "revenue": 10, "unit": "달러"},
                      {"account": "토스", "fiscal_year": 2026, "period_type": "FORECAST", "revenue": 10, "unit": "억원"}],
              sales=[{"account": "KA03", "fiscal_year": 2026, "metric": "TS", "value": 9000, "unit": "억원", "sales_rep": "TEST-REP"},
                     {"account": "KA03", "fiscal_year": 2031, "metric": "FPR", "value": 460, "unit": "억원", "sales_rep": "TEST-REP"}],
              target=[{"account": "KA03", "fiscal_year": 2031, "target_revenue": 300, "unit": "억원", "approver": "TEST-MGMT"},
                      {"account": "KA09", "fiscal_year": 2031, "target_revenue": 999999, "unit": "억원", "approver": "TEST-MGMT"}])
b1 = X.preview(f1)
rows = b1["rows"]
res["b1_summary"] = b1["summary"]
res["kb_mapped"] = next(r for r in rows if r["raw"]["account"] == "국민 은행")["mapping"]
res["kb_value_eok"] = next(r for r in rows if r["raw"]["account"] == "국민 은행")["normalized"].get("value_eok")
res["bad_unit_err"] = any("단위" in e for r in rows for e in r["errors"])
res["toss_status"] = next(r for r in rows if r["raw"]["account"] == "토스")["mapping"]["status"]
res["kakao"] = next(r for r in rows if r["raw"]["account"] == "카카오")["mapping"]
X.approve(b1["batch_id"], "TEST-APPROVER")
res["store_after_1"] = X.status()["layers"]

# 2) second upload — change KB forecast → CHANGED diff + history
f2 = workbook(tmp / "TEST_upload2.xlsx",
              actual=[{"account": "KA03", "fiscal_year": 2026, "period_type": "FORECAST", "revenue": 250, "unit": "억원"}], sales=[], target=[])
b2 = X.preview(f2)
res["b2_diff"] = b2["rows"][0]["diff"]
X.approve(b2["batch_id"], "TEST-APPROVER")
res["history_len"] = len(X.load_layer("INTERNAL_ACTUAL")["history"])

# 3) recompute (sandbox outputs)
from models.build import main as build  # noqa: E402

out = build()
acc = {a["account_id"]: a for a in out["accounts"]}
kb = acc["KA03"]
m26 = kb["m"]["2026"]
res["kb_cr"] = str(m26["CR"]["value"])
res["kb_csow_ok"] = abs(m26["CSOW"]["value"] - Decimal(250) / m26["RAW"]["value"]) < Decimal("1e-9")
res["kb_upside_ok"] = m26["UPSIDE"]["value"] == m26["FPR"]["value"] - Decimal(250)
res["kb_target"] = str(kb["m"]["target"]["TARGET"]["value"])
res["kb_target_status_note"] = kb["m"]["target"]["TARGET"]["notes"]
res["kb_recon_flags"] = {f"{r['year']}|{r['metric']}": r["flag"] for r in kb["reconciliation"]["rows"]}
res["ka09_target_error"] = any(i["code"] == "TARGET_GT_FPR" and i["severity"] == "ERROR" for i in acc["KA09"]["issues"])
res["contains_internal"] = out["meta"]["contains_internal_data"]
res["kb_na_cr"] = kb["na_class"]["2026|CR"]["class"]
res["out_dir_is_sandbox"] = str(tmp) in str(Path(os.environ["KAI_DATA_OUT"]))

# 4) API path (TestClient, same sandbox)
from fastapi.testclient import TestClient  # noqa: E402

from backend.app import app  # noqa: E402

c = TestClient(app)
with open(f2, "rb") as fh:  # same content as already-approved upload 2 → UNCHANGED
    res["api_reupload_diff"] = c.post("/api/import/preview", files={"file": ("TEST_upload2.xlsx", fh, "application/octet-stream")}).json()["rows"][0]["diff"]["type"]
f3 = workbook(tmp / "TEST_upload3.xlsx",
              actual=[{"account": "KA03", "fiscal_year": 2026, "period_type": "FORECAST", "revenue": 270, "unit": "억원"}], sales=[], target=[])
with open(f3, "rb") as fh:
    pv = c.post("/api/import/preview", files={"file": ("TEST_upload3.xlsx", fh, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}).json()
res["api_preview_diff"] = pv["rows"][0]["diff"]["type"]
ap = c.post(f"/api/import/approve/{pv['batch_id']}", json={"approver": "TEST-APPROVER"})
res["api_approve_status"] = ap.status_code
dash = c.get("/api/dashboard").json()
res["api_dash_internal"] = dash["meta"]["contains_internal_data"]
res["api_dash_kb_cr"] = next(a for a in dash["accounts"] if a["account_id"] == "KA03")["m"]["2026"]["CR"]["value"]
res["api_reject_unknown"] = c.post("/api/import/approve/NOPE", json={"approver": "x"}).status_code
print("RESULT_JSON=" + json.dumps(res, ensure_ascii=False, default=str))
